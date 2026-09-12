from __future__ import annotations

"""Attachment content extraction (added 2026-09-12).

When a file is attached to an MR we try to pull its content out so the
viewer popup can show it on screen (copy/paste-able) without downloading:

- PDF with a text layer -> page text + detected tables (pymupdf find_tables)
- scanned PDF / photos   -> tesseract OCR when the engine is installed,
  otherwise a clear "no OCR engine on this server" message
- Excel (xlsx/xlsm)      -> sheet rows
- CSV                    -> rows
- Word (.docx)           -> paragraphs + tables (zip + XML, no extra deps)

Everything is best-effort: extraction failures NEVER fail the upload.
"""

import csv
import io
import re
import zipfile
from typing import Any
from xml.etree import ElementTree

MAX_PAGES = 40
MAX_CHARS = 200_000
MAX_TABLE_ROWS = 500
MAX_SHEETS = 5

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _empty(kind: str, message: str = "", **extra: Any) -> dict[str, Any]:
    out: dict[str, Any] = {
        "ok": False,
        "kind": kind,
        "engine": "",
        "pages": 0,
        "words": 0,
        "text": "",
        "tables": [],
        "sheets": [],
        "message": message,
    }
    out.update(extra)
    return out


def _cap_text(text: str) -> str:
    text = text or ""
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n… (truncated)"
    return text


def _word_count(text: str) -> int:
    return len(re.findall(r"\S+", text or ""))


def _ocr_available() -> bool:
    import shutil

    if not shutil.which("tesseract"):
        return False
    try:
        import pytesseract  # noqa: F401

        return True
    except ImportError:
        return False


def _ocr_image_bytes(png: bytes) -> str:
    import pytesseract
    from PIL import Image

    img = Image.open(io.BytesIO(png))
    return pytesseract.image_to_string(img) or ""


def _extract_pdf(content: bytes) -> dict[str, Any]:
    try:
        import pymupdf
    except ImportError:
        return _empty("pdf", "PDF text extraction is not installed on this server (pymupdf).")
    tables: list[dict[str, Any]] = []
    text_parts: list[str] = []
    pages = 0
    has_images = False
    try:
        doc = pymupdf.open(stream=content, filetype="pdf")
        pages = min(doc.page_count, MAX_PAGES)
        for page in doc[:MAX_PAGES]:
            try:
                if page.get_images(full=True):
                    has_images = True
                text_parts.append(page.get_text("text") or "")
                try:
                    found = page.find_tables()
                    for t in found.tables:
                        rows = [[("" if c is None else str(c)).strip() for c in row] for row in t.extract]
                        rows = [r for r in rows if any(cell for cell in r)][:MAX_TABLE_ROWS]
                        if len(rows) >= 2:
                            tables.append({"page": page.number + 1, "rows": rows})
                except Exception:
                    continue  # table detection is best-effort per page
            except Exception:
                continue
        doc.close()
    except Exception as exc:
        return _empty("pdf", f"Could not read this PDF: {exc}")
    text = _cap_text("\n".join(text_parts).strip())
    if len(text) < 40 and has_images and not tables:
        if _ocr_available():
            try:
                import pymupdf

                doc = pymupdf.open(stream=content, filetype="pdf")
                ocr_parts: list[str] = []
                for page in doc[:MAX_PAGES]:
                    pix = page.get_pixmap(dpi=200)
                    ocr_parts.append(_ocr_image_bytes(pix.tobytes("png")))
                doc.close()
                text = _cap_text("\n".join(ocr_parts).strip())
                return {
                    "ok": bool(text),
                    "kind": "pdf",
                    "engine": "tesseract",
                    "pages": pages,
                    "words": _word_count(text),
                    "text": text,
                    "tables": tables,
                    "sheets": [],
                    "message": "" if text else "OCR ran but no text was recognised.",
                }
            except Exception as exc:
                return _empty("pdf", f"This PDF looks scanned and OCR failed: {exc}")
        return _empty(
            "pdf",
            "This PDF looks like a scan (no text layer). It can be viewed and downloaded, "
            "but text recognition needs the tesseract OCR engine on the server.",
            pages=pages,
        )
    return {
        "ok": bool(text or tables),
        "kind": "pdf",
        "engine": "pymupdf",
        "pages": pages,
        "words": _word_count(text),
        "text": text,
        "tables": tables,
        "sheets": [],
        "message": "" if (text or tables) else "No readable text found in this PDF.",
    }


def _extract_image(content: bytes, ext: str) -> dict[str, Any]:
    if not _ocr_available():
        return _empty(
            "image",
            "Image attached. On-screen text recognition (OCR) needs the tesseract engine on "
            "the server, which is not installed - the file can be viewed and downloaded.",
        )
    try:
        text = _cap_text(_ocr_image_bytes(content))
    except Exception as exc:
        return _empty("image", f"OCR failed on this image: {exc}")
    return {
        "ok": bool(text.strip()),
        "kind": "image",
        "engine": "tesseract",
        "pages": 1,
        "words": _word_count(text),
        "text": text,
        "tables": [],
        "sheets": [],
        "message": "" if text.strip() else "OCR ran but no text was recognised.",
    }


def _extract_sheet_rows(content: bytes) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets: dict[str, list[list[str]]] = {}
    try:
        for name in wb.sheetnames[:MAX_SHEETS]:
            ws = wb[name]
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cell.strip() for cell in cells):
                    rows.append(cells[:30])
                if len(rows) >= MAX_TABLE_ROWS:
                    break
            if rows:
                sheets[str(name)[:60]] = rows
    finally:
        wb.close()
    return sheets


def _extract_excel(content: bytes) -> dict[str, Any]:
    try:
        sheets = _extract_sheet_rows(content)
    except Exception as exc:
        return _empty("excel", f"Could not read this workbook: {exc}")
    total_rows = sum(len(v) for v in sheets.values())
    return {
        "ok": bool(sheets),
        "kind": "excel",
        "engine": "openpyxl",
        "pages": 0,
        "words": _word_count(" ".join(cell for rows in sheets.values() for row in rows for cell in row)),
        "text": "",
        "tables": [],
        "sheets": [{"name": name, "rows": rows} for name, rows in sheets.items()],
        "message": "" if sheets else "The workbook has no readable rows.",
        "table_count": total_rows,
    }


def _extract_csv(content: bytes) -> dict[str, Any]:
    try:
        text = content.decode("utf-8-sig", errors="replace")
        rows = [
            [("" if c is None else str(c)) for c in row][:30]
            for row in csv.reader(io.StringIO(text))
            if any(str(c).strip() for c in row)
        ][:MAX_TABLE_ROWS]
    except Exception as exc:
        return _empty("csv", f"Could not read this CSV: {exc}")
    flat = " ".join(cell for row in rows for cell in row)
    return {
        "ok": bool(rows),
        "kind": "csv",
        "engine": "csv",
        "pages": 0,
        "words": _word_count(flat),
        "text": "",
        "tables": rows[1:] and [{"page": 1, "rows": rows}] or [],
        "sheets": [{"name": "CSV", "rows": rows}],
        "message": "" if rows else "The CSV has no readable rows.",
        "table_count": len(rows),
    }


def _extract_docx(content: bytes) -> dict[str, Any]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            xml = zf.read("word/document.xml")
    except Exception as exc:
        return _empty("docx", "This does not look like a Word (.docx) file. Legacy .doc must be saved as .docx first.")
    try:
        root = ElementTree.fromstring(xml)
    except ElementTree.ParseError as exc:
        return _empty("docx", f"Could not parse the Word document: {exc}")

    paragraphs: list[str] = []
    tables: list[dict[str, Any]] = []
    body = root.find(f"{W_NS}body")
    elements = list(body) if body is not None else []
    for el in elements:
        if el.tag == f"{W_NS}p":
            text = "".join(t.text or "" for t in el.iter(f"{W_NS}t")).strip()
            if text:
                paragraphs.append(text)
        elif el.tag == f"{W_NS}tbl":
            rows: list[list[str]] = []
            for tr in el.findall(f"{W_NS}tr"):
                cells = []
                for tc in tr.findall(f"{W_NS}tc"):
                    cells.append(" ".join(t.text or "" for t in tc.iter(f"{W_NS}t")).strip())
                if any(c for c in cells):
                    rows.append(cells[:30])
                if len(rows) >= MAX_TABLE_ROWS:
                    break
            if len(rows) >= 2:
                tables.append({"page": 1, "rows": rows})
    text = _cap_text("\n".join(paragraphs))
    return {
        "ok": bool(text or tables),
        "kind": "docx",
        "engine": "docx-xml",
        "pages": 0,
        "words": _word_count(text),
        "text": text,
        "tables": tables,
        "sheets": [],
        "message": "" if (text or tables) else "The document appears to be empty.",
    }


def extract_attachment(filename: str, content: bytes) -> dict[str, Any]:
    """Best-effort content extraction. Never raises; never fails an upload."""
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    try:
        if ext == "pdf":
            return _extract_pdf(content)
        if ext in {"xlsx", "xlsm"}:
            return _extract_excel(content)
        if ext == "csv":
            return _extract_csv(content)
        if ext == "docx":
            return _extract_docx(content)
        if ext in {"png", "jpg", "jpeg", "webp", "gif"}:
            return _extract_image(content, ext)
        if ext == "txt":
            text = _cap_text(content.decode("utf-8", errors="replace").strip())
            return {
                "ok": bool(text),
                "kind": "txt",
                "engine": "text",
                "pages": 0,
                "words": _word_count(text),
                "text": text,
                "tables": [],
                "sheets": [],
                "message": "",
            }
        return _empty("", "No on-screen preview for this file type - it can be downloaded.")
    except Exception as exc:  # absolute last resort: the upload must still succeed
        return _empty("", f"Preview extraction failed: {exc}")
