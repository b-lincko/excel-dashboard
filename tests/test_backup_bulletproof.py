from __future__ import annotations

import shutil
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app import database  # noqa: E402
from app.config import AppConfig, save_config  # noqa: E402
from app.excel.service import ExcelService  # noqa: E402


@pytest.fixture()
def workbook(tmp_path):
    src = ROOT / "file.xlsx"
    dest = tmp_path / "file.xlsx"
    shutil.copy2(src, dest)
    cfg = AppConfig()
    cfg.excel_path = str(dest)
    cfg.backup_dir = str(tmp_path / "backups")
    save_config(cfg)
    svc = ExcelService()
    svc.invalidate()
    yield dest, svc
    cfg = AppConfig()
    cfg.excel_path = str(ROOT / "file.xlsx")
    cfg.backup_dir = str(ROOT / "backups")
    save_config(cfg)


def _mini_log(path: Path, sheet: str = "Linkco_MR_Log (SH5 & SH1)", header_row: int = 1) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    headers = ["IM Work Order #", "STATUS", "Assign to", "Supplier Name", "Required Material Details"]
    row = header_row
    for col, name in enumerate(headers, start=1):
        ws.cell(row, col).value = name
    ws.cell(row + 1, 1).value = "MR-READ-1"
    ws.cell(row + 1, 2).value = "OPEN"
    ws.cell(row + 1, 3).value = "Arun"
    ws.cell(row + 1, 4).value = "AAGE"
    ws.cell(row + 1, 5).value = "Gasket"
    wb.save(path)


def test_read_live_workbook_has_material_requests(workbook):
    dest, svc = workbook
    recs = svc.read_records_from(dest)
    assert len(recs) >= 2000
    assert any(str(r.get("work_order_id") or "").strip() for r in recs)
    assert any(str(r.get("status") or "").strip() for r in recs)


def test_read_detects_header_row_one(workbook, tmp_path):
    _, svc = workbook
    path = tmp_path / "header1.xlsx"
    _mini_log(path, header_row=1)
    recs = svc.read_records_from(path)
    assert len(recs) == 1
    assert recs[0]["work_order_id"] == "MR-READ-1"
    assert recs[0]["status"] == "OPEN"


def test_read_fuzzy_sheet_name(workbook, tmp_path):
    _, svc = workbook
    path = tmp_path / "fuzzy.xlsx"
    _mini_log(path, sheet="Linkco MR Log SH5 extra")
    recs = svc.read_records_from(path)
    assert len(recs) == 1
    assert recs[0]["work_order_id"] == "MR-READ-1"


def test_backup_still_snapshots_db_when_excel_copy_fails(workbook, monkeypatch):
    _, svc = workbook
    svc.get_all(force=True)

    def boom(*_a, **_k):
        raise OSError("simulated lock")

    monkeypatch.setattr("app.excel.service._copy_file_durable", boom)
    path = svc.create_backup(reason="manual")
    assert path is not None
    assert path.suffix.lower() == ".db"
    assert path.is_file()
    assert database.snapshot_wo_count(path) >= 1


def test_create_backup_never_raises(workbook, monkeypatch):
    _, svc = workbook

    def boom(*_a, **_k):
        raise RuntimeError("nope")

    monkeypatch.setattr("app.excel.service._copy_file_durable", boom)
    monkeypatch.setattr("app.database.snapshot_to", boom)
    path = svc.create_backup(reason="manual")
    assert path is None or path.exists()


def test_replace_from_bytes_refuses_garbage_without_wiping(workbook):
    _, svc = workbook
    svc.get_all(force=True)
    before = database.wo_cache_count()
    assert before >= 1
    with pytest.raises(ValueError, match="not an Excel"):
        svc.replace_from_bytes(b"not-an-xlsx-file" * 20, username="pytest")
    with pytest.raises(ValueError, match="old Excel"):
        svc.replace_from_bytes(b"\xd0\xcf\x11\xe0" + b"x" * 120, username="pytest")
    assert database.wo_cache_count() == before


def test_replace_from_bytes_refuses_empty_workbook(workbook, tmp_path):
    _, svc = workbook
    svc.get_all(force=True)
    before = database.wo_cache_count()
    wb = Workbook()
    wb.active.title = "Sheet1"
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="Could not find material-request"):
        svc.replace_from_bytes(path.read_bytes(), username="pytest")
    assert database.wo_cache_count() == before


def test_replace_from_bytes_refuses_tiny_workbook(workbook, tmp_path):
    dest, svc = workbook
    svc.get_all(force=True)
    before = database.wo_cache_count()
    assert before >= 10
    mini = tmp_path / "tiny.xlsx"
    _mini_log(mini)
    with pytest.raises(ValueError, match="Refusing to replace"):
        svc.replace_from_bytes(mini.read_bytes(), username="pytest")
    assert database.wo_cache_count() == before
    assert dest.is_file()
