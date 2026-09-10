from __future__ import annotations

import re
import sys
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app import database  # noqa: E402
from app.dates import week_bounds  # noqa: E402
from app.excel.service import excel_service  # noqa: E402
from app.main import app  # noqa: E402
from app.reports import period_payload, render  # noqa: E402

AS_OF = date(2026, 9, 2)  # Wednesday
SAMPLE = [
    {
        "record_id": "SH5-SH1:1",
        "work_order_id": "WO-NEW",
        "created_date": "2026-09-02",
        "status": "OPEN",
        "department": "SH5-SH1",
        "assigned_to": "Ali",
        "description": "Valve kit",
        "due_date": "2026-09-20",
        "priority": "High",
    },
    {
        "record_id": "SH5-SH1:2",
        "work_order_id": "WO-CLOSED",
        "created_date": "2026-08-20",
        "status": "CLOSED",
        "closed_date": "2026-09-02",
        "department": "SH5-SH1",
        "assigned_to": "Sara",
        "description": "Gasket set",
        "due_date": "2026-09-01",
        "priority": "Medium",
    },
    {
        "record_id": "F5:3",
        "work_order_id": "WO-OLD",
        "created_date": "2026-08-10",
        "status": "OPEN",
        "department": "F5",
        "assigned_to": "Omar",
        "description": "Pump seal",
        "due_date": "2026-08-15",
        "priority": "High",
    },
    {
        "record_id": "F5:4",
        "work_order_id": "WO-OTHER-DAY",
        "created_date": "2026-09-04",
        "status": "OPEN",
        "department": "F5",
        "assigned_to": "Ali",
        "description": "Cable tray",
        "due_date": "2026-09-30",
        "priority": "Low",
    },
    {
        "record_id": "F5:5",
        "work_order_id": "WO-OTHER-WEEK",
        "created_date": "2026-09-10",
        "status": "CLOSED",
        "closed_date": "2026-09-10",
        "department": "F5",
        "assigned_to": "Sara",
        "description": "Next week close",
        "due_date": "2026-09-12",
        "priority": "Low",
    },
]


@pytest.fixture()
def sample(monkeypatch):
    monkeypatch.setattr(excel_service, "get_all", lambda *a, **k: [dict(r) for r in SAMPLE])
    return SAMPLE


def _pdf_pages(data: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page(?!s)", data))


def test_daily_covers_chosen_day_only(sample):
    payload = period_payload("daily", {"as_of": "2026-09-02"})
    assert payload["kind"] == "daily"
    assert payload["start"] == payload["end"] == "2026-09-02"
    ids = {r["work_order_id"] for r in payload["lists"]["created"]}
    assert ids == {"WO-NEW"}
    closed = {r["work_order_id"] for r in payload["lists"]["closed"]}
    assert closed == {"WO-CLOSED"}
    overdue = {r["work_order_id"] for r in payload["lists"]["overdue"]}
    assert "WO-OLD" in overdue
    assert "WO-NEW" not in overdue
    assert "WO-OTHER-DAY" not in ids
    assert "WO-OTHER-WEEK" not in closed
    assert payload["days"] == []


def test_weekly_is_iso_week_only(sample):
    start, end = week_bounds(2026, AS_OF.isocalendar()[1])
    payload = period_payload("weekly", {"as_of": "2026-09-02"})
    assert payload["kind"] == "weekly"
    assert payload["start"] == start.date().isoformat()
    assert payload["end"] == end.date().isoformat()
    created = {r["work_order_id"] for r in payload["lists"]["created"]}
    closed = {r["work_order_id"] for r in payload["lists"]["closed"]}
    assert "WO-NEW" in created
    assert "WO-OTHER-DAY" in created  # Friday of the same ISO week
    assert "WO-OTHER-WEEK" not in created
    assert "WO-OTHER-WEEK" not in closed
    assert len(payload["days"]) == 7
    assert payload["days"][0]["name"] == "Mon"
    assert payload["days"][-1]["name"] == "Sun"


def test_daily_and_weekly_pdf_xlsx_one_page(sample):
    for kind in ("daily", "weekly"):
        pdf, name, mime = render(kind, "pdf", {"as_of": "2026-09-02"})
        assert mime == "application/pdf"
        assert pdf[:4] == b"%PDF"
        assert _pdf_pages(pdf) == 1
        assert f"{kind}_" in name

        xlsx, _, _ = render(kind, "xlsx", {"as_of": "2026-09-02"})
        wb = load_workbook(BytesIO(xlsx))
        assert wb.sheetnames == (["Daily"] if kind == "daily" else ["Weekly"])
        ws = wb.active
        assert ws.page_setup.fitToHeight == 1
        assert ws.page_setup.fitToWidth == 1


def test_reports_api_json_and_pdf(sample):
    database.init_db()
    client = TestClient(app)
    token = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"}).json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    daily = client.get("/api/reports/daily", headers=headers, params={"fmt": "json", "as_of": "2026-09-02"})
    assert daily.status_code == 200, daily.text
    body = daily.json()
    assert body["kind"] == "daily"
    assert body["start"] == "2026-09-02"
    assert body["kpis"]["created"] == 1
    pdf = client.get("/api/reports/weekly", headers=headers, params={"fmt": "pdf", "date": "2026-09-02"})
    assert pdf.status_code == 200, pdf.text
    assert pdf.content[:4] == b"%PDF"
    assert "inline" in (pdf.headers.get("content-disposition") or "")
    assert _pdf_pages(pdf.content) == 1

def test_render_supports_every_download_kind(sample):
    """Regression: render() used to return None for every kind except
    daily/weekly, so monthly/yearly/open/overdue/closed/delay/department/
    technician downloads all failed."""
    from app.reports import render

    for kind in ("monthly", "yearly", "open", "overdue", "closed", "delay", "department", "technician"):
        for fmt, sig in (("pdf", b"%PDF"), ("xlsx", b"PK"), ("csv", b"")):
            blob, name, mime = render(kind, fmt, {})
            assert blob, f"{kind}/{fmt} returned empty"
            assert blob[: len(sig)] == sig, f"{kind}/{fmt} wrong signature"
            assert kind in name, f"{kind}/{fmt} unexpected filename {name}"


def test_period_pdf_contains_day_chart(sample):
    """The weekly PDF draws a bar chart (vector drawings) and stays one page."""
    import io

    import fitz

    from app.reports import period_pdf, period_payload

    pdf = period_pdf(period_payload("weekly", {}))
    assert pdf[:4] == b"%PDF"
    doc = fitz.open(stream=pdf, filetype="pdf")
    try:
        assert doc.page_count == 1
        assert len(doc[0].get_drawings()) >= 40  # chart + tables render vectors
    finally:
        doc.close()


def test_period_xlsx_contains_native_chart(sample):
    import io

    from app.reports import period_payload, period_xlsx

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(period_xlsx(period_payload("weekly", {}))))
    ws = wb.active
    assert len(ws._charts) == 1
